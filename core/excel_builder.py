import io
from typing import List, Dict, Optional, Union
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core.normalizer import parse_data_value

class ExcelReportBuilder:
    """
    Builds professionally formatted Excel (.xlsx) workbooks from extracted data,
    applying column selection, custom headers mapping, and proper data types.
    """

    def __init__(self):
        self.header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
        self.header_font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
        self.body_font = Font(name="Aptos", size=10, color="0F172A")
        self.alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        self.white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        thin_border = Side(style="thin", color="CBD5E1")
        self.cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
        self.header_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    def build_workbook(
        self,
        df: pd.DataFrame,
        selected_columns: Optional[List[str]] = None,
        column_mapping: Optional[Dict[str, str]] = None,
        sheet_name: str = "Sheet1"
    ) -> Workbook:
        """
        Creates an openpyxl Workbook with selected columns, renamed headers, and formatting.
        - selected_columns: list of source columns to keep (in specified order)
        - column_mapping: dict of {source_col_name: target_excel_col_name}
        """
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31] # Excel limits sheet name to 31 chars

        # 1. Filter and reorder columns
        if selected_columns:
            # Only keep columns that actually exist in df
            cols_to_use = [c for c in selected_columns if c in df.columns]
            export_df = df[cols_to_use].copy()
        else:
            cols_to_use = list(df.columns)
            export_df = df.copy()

        # 2. Map header names
        mapped_headers = []
        mapping = column_mapping or {}
        for col in cols_to_use:
            mapped_headers.append(mapping.get(col, col))

        # 3. Write Header Row
        ws.row_dimensions[1].height = 26
        for col_idx, header_title in enumerate(mapped_headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=str(header_title))
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.header_border

        # 4. Write Data Rows & Apply Formats
        column_max_lens = [len(str(h)) for h in mapped_headers]

        for row_idx, (_, row) in enumerate(export_df.iterrows(), start=2):
            ws.row_dimensions[row_idx].height = 20
            row_fill = self.alt_row_fill if row_idx % 2 == 0 else self.white_fill

            for col_idx, col_name in enumerate(cols_to_use):
                raw_val = row[col_name]
                parsed_val, data_type = parse_data_value(raw_val)

                cell = ws.cell(row=row_idx, column=col_idx + 1, value=parsed_val)
                cell.font = self.body_font
                cell.fill = row_fill
                cell.border = self.cell_border

                # Alignment and number formats based on type
                if data_type == "number":
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    col_title = str(mapped_headers[col_idx]).lower()
                    is_money = any(k in col_title for k in ["เงิน", "amount", "price", "ยอด", "บาท", "discount", "total", "fee", "cost"])
                    if isinstance(parsed_val, float) or is_money:
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"
                elif data_type == "percentage":
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.00%"
                elif data_type == "date":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = "YYYY-MM-DD"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                # Track content length for auto column width
                val_len = len(str(parsed_val)) if parsed_val is not None else 0
                if val_len > column_max_lens[col_idx]:
                    column_max_lens[col_idx] = val_len

        # 5. Set Column Widths (min 12, max 50)
        for col_idx, max_len in enumerate(column_max_lens, 1):
            col_letter = get_column_letter(col_idx)
            adjusted_width = max(max_len + 4, 12)
            adjusted_width = min(adjusted_width, 50)
            ws.column_dimensions[col_letter].width = adjusted_width

        # 6. Enable Freeze Pane on top row & Auto-filter
        ws.freeze_panes = "A2"
        if export_df.shape[0] > 0 and len(cols_to_use) > 0:
            last_col_letter = get_column_letter(len(cols_to_use))
            last_row = export_df.shape[0] + 1
            ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

        return wb

    def export_to_bytes(
        self,
        df: pd.DataFrame,
        selected_columns: Optional[List[str]] = None,
        column_mapping: Optional[Dict[str, str]] = None,
        sheet_name: str = "Sheet1"
    ) -> bytes:
        """Exports the Excel workbook to an in-memory byte buffer."""
        wb = self.build_workbook(df, selected_columns, column_mapping, sheet_name)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def export_to_file(
        self,
        output_path: str,
        df: pd.DataFrame,
        selected_columns: Optional[List[str]] = None,
        column_mapping: Optional[Dict[str, str]] = None,
        sheet_name: str = "Sheet1"
    ):
        """Saves the workbook directly to disk."""
        wb = self.build_workbook(df, selected_columns, column_mapping, sheet_name)
        wb.save(output_path)
