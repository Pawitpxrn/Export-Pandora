import os
import pytest
import pandas as pd
import openpyxl
from core.normalizer import parse_data_value, clean_cell_text
from core.extractor import PDFReportExtractor
from core.excel_builder import ExcelReportBuilder
from core.profile_manager import ProfileManager

SAMPLE_PDF_PATH = os.path.join(os.path.dirname(__file__), "samples", "sample_sales_report.pdf")

def test_normalizer_values():
    # Number with commas
    val, vtype = parse_data_value("1,299,500.00")
    assert vtype == "number"
    assert val == 1299500.0

    # Negative accounting number
    val, vtype = parse_data_value("(500.00)")
    assert vtype == "number"
    assert val == -500.0

    # Integer
    val, vtype = parse_data_value("12")
    assert vtype == "number"
    assert val == 12

    # Date
    val, vtype = parse_data_value("19/09/2026")
    assert vtype == "date"
    assert str(val) == "2026-09-19"

    # Percentage
    val, vtype = parse_data_value("7.5%")
    assert vtype == "percentage"
    assert pytest.approx(val) == 0.075

def test_pdf_extractor():
    assert os.path.exists(SAMPLE_PDF_PATH), "Sample PDF should exist"
    extractor = PDFReportExtractor(SAMPLE_PDF_PATH).load_and_extract()
    
    # Assert pages and tables found
    assert extractor.pages_count >= 1
    assert len(extractor.tables) >= 1
    
    df = extractor.tables[0]
    assert len(df) >= 20 # 25 rows in sample
    assert "Item Code" in df.columns or "Product Name" in df.columns

def test_excel_builder_custom_headers():
    df = pd.DataFrame({
        "ColA": ["Item 1", "Item 2"],
        "ColB": ["1,000.00", "2,500.50"],
        "ColC": ["Ignore 1", "Ignore 2"]
    })

    builder = ExcelReportBuilder()
    # Test selecting only ColA and ColB, and renaming ColB -> "Price (THB)"
    selected = ["ColA", "ColB"]
    mapping = {"ColA": "ชื่อสินค้า", "ColB": "ราคา (บาท)"}
    
    wb = builder.build_workbook(df, selected_columns=selected, column_mapping=mapping)
    ws = wb.active

    # Check header names
    assert ws.cell(row=1, column=1).value == "ชื่อสินค้า"
    assert ws.cell(row=1, column=2).value == "ราคา (บาท)"
    assert ws.max_column == 2

    # Check values & numeric formatting
    assert ws.cell(row=2, column=1).value == "Item 1"
    assert ws.cell(row=2, column=2).value == 1000.0
    assert ws.cell(row=2, column=2).number_format == "#,##0.00"

def test_profile_manager(tmp_path):
    pm = ProfileManager(profiles_dir=str(tmp_path))
    pm.save_profile(
        profile_name="test_profile",
        selected_columns=["ColA", "ColB"],
        column_mapping={"ColA": "Product", "ColB": "Price"}
    )

    profiles = pm.list_profiles()
    assert "test_profile" in profiles

    loaded = pm.load_profile("test_profile")
    assert loaded is not None
    assert loaded["selected_columns"] == ["ColA", "ColB"]
    assert loaded["column_mapping"]["ColA"] == "Product"

def test_cashier_report_to_summary():
    user_pdf = os.path.join(os.path.dirname(__file__), "Doc", "รายงานตามเคชเชียร์.pdf")
    if os.path.exists(user_pdf):
        from core.cashier_report_parser import CashierReportParser, is_cashier_report
        assert is_cashier_report(user_pdf) is True
        parser = CashierReportParser(user_pdf).parse()
        df_summary = parser.get_payment_summary_dataframe()
        
        # Verify 3 columns as required by user's sample Excel
        assert list(df_summary.columns) == ["เลขที่ invoice", "ประเภทรับชำระ", "จำนวนเงิน"]
        assert len(df_summary) == 2
        assert df_summary.iloc[0]["เลขที่ invoice"] == "CH00126090001"
        assert df_summary.iloc[0]["ประเภทรับชำระ"] == "เงินสด"
        assert df_summary.iloc[0]["จำนวนเงิน"] == 3669.0
        assert df_summary.iloc[1]["เลขที่ invoice"] == "CH00126090002"
        assert df_summary.iloc[1]["ประเภทรับชำระ"] == "เครดิตการ์ด"
        assert df_summary.iloc[1]["จำนวนเงิน"] == 1445.0
